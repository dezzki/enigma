#!/usr/bin/env python3
"""
Enigma SE Report composer.

Reads the markdown parts in report/content/ and builds a submission-ready
Word document based on "SE Report Template.docx" (reusing its logo + styles).

Supported markdown subset:
  # / ## / ### / ####            headings
  blank line                     block separator
  - item                         bullet
  1. item                        numbered item
  | a | b |                      GitHub pipe table (with |---| separator row)
  **bold**  *italic*  `code`     inline formatting
  ```lang ... ```                fenced code block
  ![Caption](path.png)           image (own line) + caption
  <!-- PAGEBREAK -->             page break
"""
import os
import re
import sys
import subprocess
import tempfile
import zipfile

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = "/home/dezzki/Projects/enigma"
REPORT = os.path.join(ROOT, "report")
CONTENT = os.path.join(REPORT, "content")
DIAG = os.path.join(REPORT, "diagrams")
TEMPLATE = os.path.join(ROOT, "SE Report Template.docx")
OUT = os.path.join(ROOT, "Enigma-SE-Report.docx")

CHAPTER_ORDER = [
    "01-problem-statement.md",
    "02-process-model.md",
    "03-srs-a.md",
    "04-srs-b.md",
    "05-estimations.md",
    "06-scheduling.md",
    "07-risk-management.md",
    "08-design.md",
    "09-coding.md",
    "10-testing.md",
    "11-references.md",
]

ACCENT = RGBColor(0x1F, 0x3B, 0x73)

AUTHORS = [
    ("Vansh Garg", "[Roll No.]"),
    ("Natya Vidhan Biswas", "[Roll No.]"),
    ("Sarthak Maurya", "[Roll No.]"),
    ("Sehej Arora", "[Roll No.]"),
]


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #
def ensure_style(doc, name, base="Normal", stype=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        s = doc.styles.add_style(name, stype)
        try:
            s.base_style = doc.styles[base]
        except KeyError:
            pass
        return s


def strip_heading_numbering(doc):
    for lvl in range(1, 5):
        try:
            st = doc.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        pPr = st.element.find(qn("w:pPr"))
        if pPr is None:
            continue
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)


def configure_styles(doc):
    strip_heading_numbering(doc)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    specs = {
        "Heading 1": (15, True, False),
        "Heading 2": (13, True, False),
        "Heading 3": (11.5, True, False),
        "Heading 4": (11, True, True),
    }
    for name, (size, bold, italic) in specs.items():
        st = ensure_style(doc, name, base="Normal")
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = italic
        st.font.color.rgb = ACCENT
        st.paragraph_format.space_before = Pt(6 if name != "Heading 1" else 2)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.keep_with_next = True


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def extract_logo():
    with zipfile.ZipFile(TEMPLATE) as z:
        data = z.read("word/media/image1.jpeg")
    path = os.path.join(DIAG, "logo.jpeg")
    with open(path, "wb") as f:
        f.write(data)
    return path


def set_update_fields(doc):
    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")):
        settings.remove(el)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_runs(paragraph, text):
    tokens = re.split(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif (
            tok.startswith("*")
            and tok.endswith("*")
            and len(tok) > 2
            and not tok.startswith("**")
        ):
            paragraph.add_run(tok[1:-1]).italic = True
        else:
            paragraph.add_run(tok)


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, text)
    return p


def add_heading(doc, text, level):
    level = max(1, min(level, 4))
    p = doc.add_paragraph(style=f"Heading {level}")
    add_runs(p, text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.35)
    pf.first_line_indent = Inches(-0.22)
    pf.space_after = Pt(3)
    add_runs(p, "\u2022  " + text)
    return p


def add_numbered(doc, text, n):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.42)
    pf.first_line_indent = Inches(-0.3)
    pf.space_after = Pt(3)
    add_runs(p, f"{n}.  " + text)
    return p


def add_code_block(doc, code):
    lines = code.split("\n")
    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(6 if idx == 0 else 0)
        pf.line_spacing = 1.0
        pf.left_indent = Inches(0.15)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        shade(p, "F4F6F8")


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


def set_orientation(section, landscape):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        w, h = section.page_height, section.page_width
        if w > h:
            w, h = h, w
        section.page_width = max(w, h)
        section.page_height = min(w, h)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        w, h = section.page_width, section.page_height
        section.page_width = min(w, h)
        section.page_height = max(w, h)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)


def add_image(doc, path, caption, state=None):
    path = path.replace("\\", "/")
    if not os.path.isabs(path):
        path = os.path.join(REPORT, path)
    if not os.path.exists(path):
        p = doc.add_paragraph()
        add_runs(p, f"[missing figure: {os.path.basename(path)}]")
        return
    size = png_size(path)
    if not size:
        doc.add_picture(path, width=Inches(6.3))
        return
    px_w, px_h = size
    nat_w, nat_h = px_w / 96.0, px_h / 96.0

    landscape = nat_w > 9.0
    if landscape:
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        set_orientation(sec, True)
        max_w, max_h = 9.5, 6.4
    else:
        max_w, max_h = 6.8, 7.0

    scale = min(max_w / nat_w, max_h / nat_h, 1.0)
    draw_w = nat_w * scale

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(path, width=Inches(draw_w))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(4)
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    if landscape:
        nxt = doc.add_section(WD_SECTION.NEW_PAGE)
        set_orientation(nxt, False)


def parse_table_row(row):
    row = row.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [c.strip() for c in row.split("|")]


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells) and cells


def set_table_cell_margins(table, top=10, bottom=10, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("w:top", top), ("w:left", left), ("w:bottom", bottom), ("w:right", right)):
        el = OxmlElement(tag)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def add_table(doc, raw_rows):
    rows = [parse_table_row(r) for r in raw_rows]
    rows = [r for r in rows if not is_separator_row(r)]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    set_table_cell_margins(table)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(ncols):
            txt = row[j] if j < len(row) else ""
            para = cells[j].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            add_runs(para, txt)
            for run in para.runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.bold = True
    return table


# --------------------------------------------------------------------------- #
# markdown processor
# --------------------------------------------------------------------------- #
def process_markdown(doc, path, state):
    if not os.path.exists(path):
        p = doc.add_paragraph()
        add_runs(p, f"[missing part: {os.path.basename(path)}]")
        return
    with open(path, encoding="utf-8") as f:
        lines = f.read().replace("\r\n", "\n").split("\n")

    i = 0
    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        if s == "":
            i += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            i += 1
            continue
        if s.startswith(">"):
            s = s.lstrip("> ").strip()
            if not s:
                i += 1
                continue
        if s.startswith("<!--") and "PAGEBREAK" in s.upper():
            doc.add_page_break()
            i += 1
            continue
        if s.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, "\n".join(buf))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            started = state["started"]
            if level == 1:
                state["started"] = True
            p = add_heading(doc, text, level)
            if level == 1 and started:
                p.paragraph_format.page_break_before = True
            i += 1
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", s)
        if m:
            add_image(doc, m.group(2), m.group(1), state)
            i += 1
            continue
        if s.startswith("|") and s.endswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            add_table(doc, block)
            continue
        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            while i < len(lines):
                mm = re.match(r"^[-*]\s+(.*)$", lines[i].strip())
                if not mm:
                    break
                add_bullet(doc, mm.group(1))
                i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            while i < len(lines):
                mm = re.match(r"^(\d+)\.\s+(.*)$", lines[i].strip())
                if not mm:
                    break
                add_numbered(doc, mm.group(2), int(mm.group(1)))
                i += 1
            continue
        # paragraph
        buf = [s]
        i += 1
        while i < len(lines):
            ns = lines[i].strip()
            if (
                ns == ""
                or ns.startswith("#")
                or ns.startswith("|")
                or ns.startswith("```")
                or ns.startswith("![")
                or ns.startswith("<!--")
                or re.match(r"^[-*]\s+", ns)
                or re.match(r"^\d+\.\s+", ns)
            ):
                break
            buf.append(ns)
            i += 1
        add_paragraph(doc, " ".join(buf))


# --------------------------------------------------------------------------- #
# front matter
# --------------------------------------------------------------------------- #
def front_heading(doc, text, size=18, page_break_before=False):
    p = doc.add_paragraph()
    if page_break_before:
        p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = ACCENT
    return p


def centered(doc, text, size=11, bold=False, italic=False, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def add_title_page(doc, logo):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(logo, width=Inches(1.4))
    doc.add_paragraph()
    centered(doc, "University of Delhi", 15, bold=True, space_after=2)
    centered(doc, "Keshav Mahavidyalaya", 14, space_after=2)
    centered(doc, "Department of Computer Science", 13, space_after=18)
    centered(doc, "Software Engineering", 13, italic=True, space_after=2)
    centered(doc, "Project Report", 13, italic=True, space_after=24)
    centered(doc, "Enigma", 30, bold=True, space_after=6)
    centered(doc, "Social Challenges with Real Outcomes", 14, italic=True, space_after=30)

    tbl = doc.add_table(rows=0, cols=2)
    tbl.autofit = True
    head = tbl.add_row().cells
    for cell, txt in zip(head, ["Submitted by:", "Supervisor:"]):
        r = cell.paragraphs[0].add_run(txt)
        r.bold = True
    for name, roll in AUTHORS:
        cells = tbl.add_row().cells
        cells[0].paragraphs[0].add_run(f"{name}  ({roll})")
        cells[1].paragraphs[0].add_run("Prof. Anjali Thukral" if name == AUTHORS[0][0] else "")
    doc.add_paragraph()
    centered(doc, "2026", 13, bold=True, space_after=4)


def add_certificate(doc):
    doc.add_paragraph()
    front_heading(doc, "CERTIFICATE", 18, page_break_before=True)
    doc.add_paragraph()
    body = (
        "This is to certify that the work embodied in this project entitled "
        "\u201cEnigma \u2014 Social Challenges with Real Outcomes\u201d has been "
        "carried out in the Department of Computer Science, Keshav Mahavidyalaya, "
        "University of Delhi. The work presented in this project is original and has "
        "not been submitted in part or full to this or any other University for any "
        "diploma or degree."
    )
    p = add_paragraph(doc, body)
    p.paragraph_format.line_spacing = 1.5
    for _ in range(6):
        doc.add_paragraph()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.autofit = True
    tbl.rows[0].cells[0].paragraphs[0].add_run("_______________________").bold = False
    tbl.rows[0].cells[1].paragraphs[0].add_run("_______________________").bold = False
    tbl.rows[1].cells[0].paragraphs[0].add_run("Candidate").italic = True
    tbl.rows[1].cells[1].paragraphs[0].add_run("Supervisor").italic = True
    for row in tbl.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_static_toc(doc, entries, pages, page_break_before=False):
    front_heading(doc, "Table of Contents", 18, page_break_before=page_break_before)
    for i, (level, text) in enumerate(entries):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.space_before = Pt(6 if level == 1 else 0)
        pf.line_spacing = 1.0
        pf.left_indent = Inches(0.0 if level == 1 else (0.28 if level == 2 else 0.56))
        pf.tab_stops.add_tab_stop(
            Inches(6.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        clean = re.sub(r"[`*]", "", text)
        r = p.add_run(clean)
        r.bold = level == 1
        r.font.size = Pt(10.5)
        number = "" if pages[i] in (None, "") else str(pages[i])
        r2 = p.add_run("\t" + number)
        r2.font.size = Pt(10.5)


def collect_headings():
    entries = []
    for name in CHAPTER_ORDER:
        path = os.path.join(CONTENT, name)
        if not os.path.exists(path):
            continue
        in_fence = False
        with open(path, encoding="utf-8") as f:
            for raw in f:
                s = raw.strip()
                if s.startswith("```"):
                    in_fence = not in_fence
                    continue
                if in_fence:
                    continue
                m = re.match(r"^(#{1,2})\s+(.*)$", s)
                if m:
                    entries.append((len(m.group(1)), m.group(2).strip()))
    return entries


def _norm_heading(s):
    s = re.sub(r"[`*]", "", s)
    return re.sub(r"\s+", "", s).lower()


def pdf_outline(pdf_path):
    from pypdf import PdfReader

    reader = PdfReader(pdf_path)
    out = []

    def walk(items, depth):
        for it in items:
            if isinstance(it, list):
                walk(it, depth + 1)
            else:
                try:
                    pg = reader.get_destination_page_number(it) + 1
                except Exception:
                    pg = None
                out.append((depth, getattr(it, "title", ""), pg))

    walk(reader.outline, 1)
    return out


def map_pages(entries, outline):
    pages = [None] * len(entries)
    # primary: align outline (in document order) with markdown headings
    if len(outline) == len(entries):
        return [pg for _, _, pg in outline]
    cur = 0
    for _, title, pg in outline:
        t = _norm_heading(title)
        j = cur
        while j < len(entries) and _norm_heading(entries[j][1]) != t:
            j += 1
        if j < len(entries):
            pages[j] = pg
            cur = j + 1
    return pages


def add_toc(doc):
    """Deprecated: kept for reference; static TOC is generated instead."""
    raise NotImplementedError("use add_static_toc")


def add_footer_page_numbers(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(it); run._r.append(e)
    for sec in doc.sections[1:]:
        sec.footer.is_linked_to_previous = True
        sec.different_first_page_header_footer = False


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def build_document(entries, pages):
    doc = Document(TEMPLATE)
    clear_body(doc)
    configure_styles(doc)
    logo = extract_logo()

    doc.core_properties.title = "Enigma — Social Challenges with Real Outcomes"
    doc.core_properties.author = ", ".join(a for a, _ in AUTHORS)

    add_title_page(doc, logo)
    add_certificate(doc)

    front_heading(doc, "Abstract", 18, page_break_before=True)
    process_markdown(doc, os.path.join(CONTENT, "00-abstract.md"), {"started": True})

    add_static_toc(doc, entries, pages, page_break_before=True)

    state = {"started": True}
    for name in CHAPTER_ORDER:
        process_markdown(doc, os.path.join(CONTENT, name), state)

    add_footer_page_numbers(doc)
    set_update_fields(doc)
    return doc


def convert_to_pdf(docx_path, outdir):
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=600,
    )
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def main():
    entries = collect_headings()
    if not entries:
        print("ERROR: no headings found; aborting.")
        sys.exit(1)

    work = tempfile.mkdtemp(prefix="enigma_build_")
    # pass 1 — placeholder TOC (fixed line count) to establish pagination
    doc1 = build_document(entries, [None] * len(entries))
    pass1 = os.path.join(work, "pass1.docx")
    doc1.save(pass1)
    pdf1 = convert_to_pdf(pass1, work)
    outline = pdf_outline(pdf1)
    pages = map_pages(entries, outline)
    if sum(p is not None for p in pages) < len(pages) * 0.9:
        print("WARNING: could not map all heading pages; check TOC accuracy")

    # pass 2 — real page numbers
    doc2 = build_document(entries, pages)
    doc2.save(OUT)

    missing = [
        n
        for n in CHAPTER_ORDER + ["00-abstract.md"]
        if not os.path.exists(os.path.join(CONTENT, n))
    ]
    print(f"Wrote {OUT}")
    if missing:
        print("WARNING missing parts:", ", ".join(missing))
    print(
        f"paragraphs={len(doc2.paragraphs)} tables={len(doc2.tables)} "
        f"sections={len(doc2.sections)} toc_entries={len(entries)} "
        f"mapped={sum(p is not None for p in pages)}"
    )


if __name__ == "__main__":
    main()
